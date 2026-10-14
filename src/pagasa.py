"""
Version 4, step 7. PAGASA's own record against IBTrACS.

Inputs (PAGASA's files are not redistributed; they are available from PAGASA)
    list    Climate Division list of tropical cyclones that crossed the PAR,
            1960-2024 (a spreadsheet; the list is on its sheet "List of TC"):
            local name, PAR entry and exit dates, estimated maximum sustained
            wind, category, and a mark for storms that entered more than once
    tracks  Weather Division 6-hourly positions and categories, 1948-2026
            (a CSV file). Its own note says it is preliminary, based
            mostly on the JTWC best track before the 1980s, on the JMA best track
            from the 1980s to about 2009, and on PAGASA's records after that.
    IBTrACS ibtracs.WP.list.v04r01.csv, and data/storm_flags.csv from threshold.py

Steps
    1  clean the 6-hourly file. A track is a contiguous block of rows with the
       same name. Timestamps that break the 6-hourly sequence are repaired:
       January fixes that carry the previous year; a single fix out of sequence
       between two fixes 12 h apart, or at either end of a track; and blocks of
       fixes whose month or year is off, shifted by the whole months or years
       that make the track continuous. A block that no such shift connects is
       split off as its own track. Exact repeats of a time are then dropped.
    2  link each listed storm to its tracks: same name first, then a one-letter
       spelling difference or a TS/TY prefix on a number name, with a fix inside
       the PAR within 3 days of the listed entry and exit
    3  match each listed storm (all its linked fixes) to an IBTrACS SID, one to
       one: candidate pairs share at least 3 times (all of a shorter track), and
       are taken in order of mean separation at those times, a spur charged
       25 km against a main track; separations above MAXKM (200 km) are refused,
       and 400 km is kept as a sensitivity
    4  PAGASA's annual series 1960-2023, and how the list and the IBTrACS PAR
       storm classes (decompose.py, main track) map onto each other
    5  the trend test of multiverse.fit for start years 1960-2000, calibrated
       p-values under fitted noise as in calibrate.py, and the step models of
       dating.py, for PAGASA's series and for IBTrACS over the same years

Intensity inside the PAR comes from the 6-hourly categories (TD; TS, STS and the
TSS used in 1989-2000; TY, STY and the TYS used in 1989-2000), as for IBTrACS.
The listed category is kept as a sensitivity series.

Outputs (in OUTDIR; derived from PAGASA's data, so keep them out of the public repository)
    pagasa_storms.csv       one row per listed storm: link, IBTrACS match, classes
    pagasa_ibtracs_par.csv  IBTrACS PAR main-track storms 1960-2023: class, and whether listed
    pagasa_tracks_log.csv   cleaning and matching counts
    pagasa_annual.csv       annual series 1960-2023
    pagasa_trends.csv       slope and verdict per series and start year
    pagasa_calibrated_p.csv two-sided p-values under fitted noise (blank for a boundary fit)
    pagasa_dating.csv       selected step or smooth model per series

Usage
    python src/pagasa.py LIST.xlsx TRACKS.csv ibtracs.WP.list.v04r01.csv data/storm_flags.csv OUTDIR
"""
import os
import sys
import warnings

import numpy as np
import openpyxl
import pandas as pd
from matplotlib.path import Path

from calibrate import arma_sim, detrended, fitted_models
from classes import annual as ib_annual, fit_many, load_flags
from dating import Binomial, Poisson, date_series
from decompose import classes as ib_classes
from multiverse import fit

warnings.filterwarnings("ignore")

PAR = Path([(115, 5), (115, 15), (120, 21), (120, 25), (135, 25), (135, 5)])
Y0, Y1 = 1960, 2023
YEARS = list(range(Y0, Y1 + 1))
STARTS = range(1960, 2001)
MAXKM, PREFER_MAIN_KM = 200.0, 25.0
RANK = {"TD": 1, "TS": 2, "TSS": 2, "STS": 2, "TY": 3, "TYS": 3, "STY": 3}
PARTITION = ("ts_or_stronger", "depression_only", "below_depression", "no_intensity")


def in_par(lon, lat):
    return PAR.contains_points(np.c_[lon, lat], radius=-1e-9)


# ---------------------------------------------------------------- step 1
def _good(tv):
    dt = np.diff(tv).astype("timedelta64[h]").astype(int)
    return (dt > 0) & (dt <= 48)


def _repair(tt, log):
    """Repair one track's timestamps (a list of Timestamps in file order).
    Returns the repaired list and the positions where the track must be split."""
    n = len(tt)
    for j in range(1, n):                                    # January with the previous year
        if tt[j] < tt[j - 1] and tt[j - 1].month == 12 and tt[j].month == 1 and tt[j].year == tt[j - 1].year:
            y = tt[j].year
            for m in range(j, n):
                if tt[m].month == 1 and tt[m].year == y:
                    tt[m] += pd.DateOffset(years=1)
                    log["year_fixed"] += 1
    six, twelve = pd.Timedelta(hours=6), pd.Timedelta(hours=12)
    for j in range(1, n - 1):                                # one fix out of sequence
        if not (tt[j - 1] < tt[j] < tt[j + 1]) and tt[j + 1] - tt[j - 1] == twelve:
            tt[j] = tt[j - 1] + six
            log["single_fixed"] += 1
    if n >= 3 and tt[-1] <= tt[-2] and tt[-2] - tt[-3] == six:   # last fix
        tt[-1] = tt[-2] + six
        log["single_fixed"] += 1
    if n >= 3 and tt[0] >= tt[1] and tt[2] - tt[1] == six:       # first fix
        tt[0] = tt[1] - six
        log["single_fixed"] += 1
    splits = []
    for _ in range(4):                                       # whole months or years off
        tv = np.array(tt, dtype="datetime64[ns]")
        ok = _good(tv)
        if ok.all():
            break
        cuts = [0, *(np.flatnonzero(~ok) + 1), n]
        blocks = [(a, b) for a, b in zip(cuts[:-1], cuts[1:])]
        main = max(range(len(blocks)), key=lambda i: blocks[i][1] - blocks[i][0])
        changed = False
        for i, (a, b) in enumerate(blocks):
            if i == main:
                continue
            for off in [pd.DateOffset(months=k) for k in (-1, 1, -2, 2)] + [pd.DateOffset(years=k) for k in (-1, 1)]:
                trial = list(tt)
                for m in range(a, b):
                    trial[m] = tt[m] + off
                tv2 = np.array(trial, dtype="datetime64[ns]")
                lo, hi = max(a - 1, 0), min(b, n - 1)
                if _good(tv2[lo:hi + 1]).all():
                    tt = trial
                    log["block_shifted"] += int(b - a)
                    changed = True
                    break
            if changed:
                break
        if not changed:
            break
    tv = np.array(tt, dtype="datetime64[ns]")
    dh = np.diff(tv).astype("timedelta64[h]").astype(int)
    splits = list(np.flatnonzero((dh <= 0) | (dh > 30 * 24)) + 1)   # what no shift connects
    return tt, splits


def clean_tracks(path):
    raw = pd.read_csv(path, dtype=str, keep_default_na=False).iloc[:, :6]
    raw = raw[(raw != "").any(axis=1)].reset_index(drop=True)
    d = pd.DataFrame({
        "name": raw.tc_name_local.str.strip().str.upper(),
        "intl": raw.tc_name_int.str.strip().str.upper(),
        "cat": raw.tc_cat.str.strip().str.upper(),
        "t": pd.to_datetime(raw.tc_date, format="%Y-%m-%d %H"),
        "lon": raw.lon.astype(float), "lat": raw.lat.astype(float)})
    d["base"] = d.name.str.replace("(REISSUE)", "", regex=False).str.strip()
    blk = (d.name != d.name.shift()).cumsum().to_numpy()
    log = {"rows": len(d), "year_fixed": 0, "single_fixed": 0, "block_shifted": 0, "split_off": 0}
    t = list(d.t)
    track = np.zeros(len(d), dtype=int)
    nxt = 0
    for k in np.unique(blk):
        idx = np.flatnonzero(blk == k)
        tt, splits = _repair([t[i] for i in idx], log)
        for j, i in enumerate(idx):
            t[i] = tt[j]
        bounds = [0, *splits, len(idx)]
        for a, b in zip(bounds[:-1], bounds[1:]):
            track[idx[a:b]] = nxt
            nxt += 1
        log["split_off"] += len(splits)
    d["t"] = pd.to_datetime(t)
    d["track"] = track
    d = d.sort_values(["track", "t"], kind="stable")
    dup = d.duplicated(["track", "t"])
    log["repeats_dropped"] = int(dup.sum())
    d = d[~dup].reset_index(drop=True)
    gaps = d.groupby("track").t.apply(lambda s: s.diff().max())
    log["tracks"] = int(d.track.nunique())
    log["tracks_with_gap_over_48h"] = int((gaps > pd.Timedelta(hours=48)).sum())
    d["inpar"] = in_par(d.lon.to_numpy(), d.lat.to_numpy())
    d["rank"] = d.cat.map(RANK).fillna(0).astype(int)
    return d, log


# ---------------------------------------------------------------- step 2
def read_list(path):
    ws = openpyxl.load_workbook(path, data_only=True)["List of TC"]
    rows, r = [], 5
    while isinstance(ws.cell(r, 2).value, (int, float)):
        v = [ws.cell(r, c).value for c in range(2, 10)]
        rows.append(dict(YEAR=int(v[0]), MONTH=int(v[1]), TYPE=str(v[2]).strip(), NAME=str(v[3]).strip().upper(),
                         BEG=pd.Timestamp(v[4]), END=pd.Timestamp(v[5]), MSW=float(v[6]), MULTI=(v[7] == "**")))
        r += 1
    L = pd.DataFrame(rows)
    L.index.name = "list_id"
    return L


def _lev1(a, b):
    if abs(len(a) - len(b)) > 1:
        return False
    if len(a) == len(b):
        return sum(x != y for x, y in zip(a, b)) == 1
    s, l = (a, b) if len(a) < len(b) else (b, a)
    return any(l[:i] + l[i + 1:] == s for i in range(len(l)))


def _near_name(a, b):
    if a == b:
        return False
    if a[2:].isdigit() and a[2:] == b[2:]:
        return True
    return _lev1(a, b)


def link(L, d):
    tr = d.groupby("track").agg(base=("base", "first"), t0=("t", "min"), t1=("t", "max"))
    ip = d[d.inpar].groupby("track").t.agg(["min", "max"]).rename(columns={"min": "p0", "max": "p1"})
    tr = tr.join(ip)
    day = pd.Timedelta(days=3)
    out = []
    for i, s in L.iterrows():
        exact = tr[tr.base == s.NAME]
        near = tr[tr.base.apply(lambda b: _near_name(b, s.NAME))]
        hit, how, nm = pd.DataFrame(), None, None
        for cand, kind in ((exact, "exact"), (near, "near")):
            h = cand[cand.p0.notna() & (cand.p0 <= s.END + day) & (cand.p1 >= s.BEG - day)]
            if len(h):
                hit, how, nm = h, "inside PAR", kind
                break
        if hit.empty:
            for cand, kind in ((exact, "exact"), (near, "near")):
                h = cand[(cand.t0 <= s.END + day) & (cand.t1 >= s.BEG - day)]
                if len(h):
                    hit, how, nm = h, "track never inside PAR", kind
                    break
        if hit.empty:
            h = exact[(exact.t0.dt.year == s.YEAR) & exact.p0.notna()]
            if len(h) == 1:
                hit, how, nm = h, "same year, dates disagree", "exact"
        out.append(dict(list_id=i, tracks=" ".join(str(k) for k in hit.index), link=how, name_match=nm))
    return pd.DataFrame(out).set_index("list_id")


# ---------------------------------------------------------------- step 3
def load_ibtracs(path):
    cols = ["SID", "SEASON", "ISO_TIME", "LAT", "LON", "TRACK_TYPE", "NAME"]
    ib = pd.read_csv(path, usecols=cols, skiprows=[1], low_memory=False, keep_default_na=False)
    ib = ib[pd.to_numeric(ib.SEASON, errors="coerce") >= 1945].copy()
    ib["t"] = pd.to_datetime(ib.ISO_TIME)
    ib["lat"] = pd.to_numeric(ib.LAT, errors="coerce")
    ib["lon"] = pd.to_numeric(ib.LON, errors="coerce")
    return ib.dropna(subset=["lat", "lon"])


def _km(lat1, lon1, lat2, lon2):
    p1, p2 = np.radians(lat1), np.radians(lat2)
    dl = np.radians(lon2 - lon1)
    a = np.sin((p2 - p1) / 2) ** 2 + np.cos(p1) * np.cos(p2) * np.sin(dl / 2) ** 2
    return 2 * 6371.0 * np.arcsin(np.sqrt(np.clip(a, 0, 1)))


def candidates(groups, ib):
    """Every (PAGASA key, IBTrACS SID) pair with enough common times, and its mean separation."""
    m = groups[["key", "t", "lat", "lon"]].merge(ib[["SID", "t", "lat", "lon", "TRACK_TYPE"]], on="t",
                                                 suffixes=("", "_ib"))
    m["km"] = _km(m.lat.to_numpy(), m.lon.to_numpy(), m.lat_ib.to_numpy(), m.lon_ib.to_numpy())
    agg = m.groupby(["key", "SID"]).agg(n=("km", "size"), mean_km=("km", "mean"),
                                        track_type=("TRACK_TYPE", "first")).reset_index()
    nfix = groups.groupby("key").size()
    agg["need"] = agg.key.map(nfix).clip(upper=3)
    return agg[agg.n >= agg.need].drop(columns="need")


def assign(cand, maxkm):
    """One-to-one: pairs taken in order of mean separation, a spur charged PREFER_MAIN_KM,
    each PAGASA storm and each SID used once, separations above maxkm refused."""
    c = cand.assign(cost=cand.mean_km + np.where(cand.track_type == "main", 0.0, PREFER_MAIN_KM))
    c = c[c.mean_km <= maxkm].sort_values(["cost", "n"], ascending=[True, False])
    seen_k, seen_s, rows = set(), set(), []
    for r in c.itertuples():
        if r.key in seen_k or r.SID in seen_s:
            continue
        seen_k.add(r.key)
        seen_s.add(r.SID)
        rows.append(dict(key=r.key, SID=r.SID, track_type=r.track_type, n_common=int(r.n), mean_km=float(r.mean_km)))
    return pd.DataFrame(rows, columns=["key", "SID", "track_type", "n_common", "mean_km"]).set_index("key")


# ---------------------------------------------------------------- steps 4 and 5
def series_of(years, values=None):
    s = pd.Series(years).value_counts() if values is None else values
    return s.reindex(YEARS, fill_value=0).astype(int)


def main():
    lst, trk, ibp, flags, out = sys.argv[1:6]
    os.makedirs(out, exist_ok=True)
    d, log = clean_tracks(trk)
    L = read_list(lst)
    lk = link(L, d)
    S = L.join(lk)
    used = [k for s in S.tracks for k in s.split()]
    log["list_storms"] = len(L)
    log["list_linked_inside_par"] = int((S.link == "inside PAR").sum())
    log["list_linked_track_outside_par"] = int((S.link == "track never inside PAR").sum())
    log["list_linked_dates_disagree"] = int((S.link == "same year, dates disagree").sum())
    log["list_unlinked"] = int(S.link.isna().sum())
    log["list_linked_by_near_name"] = int((S.name_match == "near").sum())
    log["tracks_linked_twice"] = int(pd.Series(used).duplicated().sum())

    # PAGASA intensity inside the PAR, from the linked fixes
    fx = d.assign(track=d.track.astype(str))
    rank_in, n_in = {}, {}
    for i, s in S.iterrows():
        f = fx[fx.track.isin(s.tracks.split())]
        fin = f[f.inpar]
        rank_in[i] = int(fin["rank"].max()) if len(fin) else np.nan
        n_in[i] = len(fin)
    S["rank_in_par"] = pd.Series(rank_in)
    S["fixes_in_par"] = pd.Series(n_in)
    listed_rank = S.TYPE.map({"TD": 1, "TS": 2, "STS": 2, "TY": 3, "STY": 3})
    S["ts_6h"] = np.where(S.rank_in_par.notna(), S.rank_in_par >= 2, listed_rank >= 2)
    S["ts_listed"] = listed_rank >= 2

    # match listed storms (all their linked fixes) to IBTrACS, one to one
    ib = load_ibtracs(ibp)
    g = pd.concat([fx[fx.track.isin(s.tracks.split())].assign(key=i) for i, s in S.iterrows() if s.tracks])
    cand = candidates(g, ib)
    S["n_candidates"] = cand.groupby("key").size().reindex(S.index).fillna(0).astype(int)
    S["nearest_km"] = cand.groupby("key").mean_km.min().reindex(S.index)
    S = S.join(assign(cand, MAXKM))
    alt = assign(cand, 2 * MAXKM)
    S["SID_400"] = alt.SID.reindex(S.index)
    why = {}
    for i in S.index[S.SID.isna()]:
        c = cand[cand.key == i]
        if c.empty:
            why[i] = "no IBTrACS storm at common times"
        elif (c.mean_km <= MAXKM).any():
            why[i] = "its IBTrACS storm is matched to another listed storm"
        elif (c.mean_km <= 2 * MAXKM).any():
            why[i] = "nearest IBTrACS storm 200-400 km"
        else:
            why[i] = "nearest IBTrACS storm beyond 400 km"
    S["unmatched_reason"] = pd.Series(why, dtype=object)
    log["list_with_candidate"] = int((S.n_candidates > 0).sum())
    log["list_matched"] = int(S.SID.notna().sum())
    log["list_matched_400km"] = int(S.SID_400.notna().sum())
    log["list_matched_spur"] = int((S.SID.notna() & (S.track_type != "main")).sum())

    # IBTrACS PAR main-track storms and their classes
    fl = load_flags(flags, "PAR")
    cl = ib_classes(fl)
    cls = pd.Series("", index=fl.index)
    for k in PARTITION:
        cls[cl[k]] = k
    S["ib_in_par_main"] = S.SID.isin(fl.index)
    S["ib_class"] = S.SID.map(cls)
    S["ib_season"] = S.SID.map(fl.SEASON)
    win = fl[fl.SEASON.between(Y0, Y1)]
    U = pd.DataFrame({"SEASON": win.SEASON, "ib_class": cls[win.index],
                      "in_list": win.index.isin(set(S.SID.dropna())),
                      "in_list_400km": win.index.isin(set(S.SID_400.dropna()))})
    U.to_csv(os.path.join(out, "pagasa_ibtracs_par.csv"))
    S.to_csv(os.path.join(out, "pagasa_storms.csv"))

    # annual series 1960-2023 (PAGASA by listed year, IBTrACS by season)
    SL = S[S.YEAR.between(Y0, Y1)]
    A = pd.DataFrame(index=pd.Index(YEARS, name="year"))
    A["pagasa_all"] = series_of(SL.YEAR)
    A["pagasa_ts"] = series_of(SL.YEAR[SL.ts_6h])
    A["pagasa_td"] = A.pagasa_all - A.pagasa_ts
    A["pagasa_ts_listed"] = series_of(SL.YEAR[SL.ts_listed])
    for k in ("total", *PARTITION):
        A[f"ib_{k}"] = ib_annual(fl, cl[k]).reindex(YEARS).astype(int)
    for k in PARTITION:
        A[f"ib_{k}_listed"] = series_of(U.SEASON[(U.ib_class == k) & U.in_list])
    A["ib_listed"] = series_of(U.SEASON[U.in_list])
    A.to_csv(os.path.join(out, "pagasa_annual.csv"))

    # trend tests
    TS = ("pagasa_all", "pagasa_ts", "pagasa_td", "pagasa_ts_listed", "ib_total", "ib_ts_or_stronger",
          "ib_depression_only", "ib_no_intensity")
    tr = []
    for name in TS:
        for y in STARTS:
            sl, lo, hi, vd = fit(A[name].loc[y:Y1])
            tr.append(dict(series=name, start=y, slope=sl, lo=lo, hi=hi, verdict=vd))
    pd.DataFrame(tr).to_csv(os.path.join(out, "pagasa_trends.csv"), index=False)

    cp = []
    for name in ("pagasa_all", "pagasa_ts", "pagasa_td", "ib_total", "ib_ts_or_stronger"):
        models = fitted_models(detrended(A[name].to_numpy(float)))
        for y0 in (1960, 1977):
            vv = A[name].loc[y0:Y1].to_numpy(float)
            xx = np.arange(len(vv), dtype=float)
            sl, band, vd = fit_many(vv[None, :], xx)
            for lab, (ar, ma, sig) in models.items():
                # an ARMA fit on the invertibility or stationarity boundary has no usable spectrum
                boundary = bool(np.any(np.abs(ar) >= 0.99) or np.any(np.abs(ma) >= 0.99))
                Z = arma_sim(ar, ma, sig, len(vv), 20000, np.random.default_rng(1900 + y0))
                p = float(np.mean(np.abs(fit_many(Z, xx)[0]) >= abs(sl[0])))
                cp.append(dict(series=name, start=y0, slope=sl[0], test_verdict={-1: "-", 0: "0", 1: "+"}[int(vd[0])],
                               model=lab.split(" ")[0] if "AIC" in lab else lab,
                               model_role="AIC" if "AIC" in lab else "ARMA(1,1)",
                               params="ar=" + ",".join(f"{c:.2f}" for c in ar) + " ma=" + ",".join(f"{c:.2f}" for c in ma),
                               boundary_fit=boundary, p_two_sided=np.nan if boundary else p))
    pd.DataFrame(cp).to_csv(os.path.join(out, "pagasa_calibrated_p.csv"), index=False)

    t = np.array(YEARS)
    dt = []
    for name in ("pagasa_all", "pagasa_ts", "pagasa_td", "ib_total", "ib_ts_or_stronger", "ib_depression_only",
                 "ib_no_intensity"):
        dt.append(dict(series=name, family="count", **date_series(Poisson(A[name].to_numpy(float)), t)))
    for name, k, n in (("listed share of IBTrACS total", "ib_listed", "ib_total"),
                       ("listed share of IBTrACS TS or stronger", "ib_ts_or_stronger_listed", "ib_ts_or_stronger"),
                       ("TS share of PAGASA list", "pagasa_ts", "pagasa_all")):
        dt.append(dict(series=name, family="share",
                       **date_series(Binomial(A[k].to_numpy(float), A[n].to_numpy(float)), t)))
    pd.DataFrame(dt).to_csv(os.path.join(out, "pagasa_dating.csv"), index=False)
    pd.Series(log).to_csv(os.path.join(out, "pagasa_tracks_log.csv"), header=["value"])
    for k, v in log.items():
        print(f"{k:32s} {v}")
    print(f"written: pagasa_* in {out}/")


if __name__ == "__main__":
    main()
