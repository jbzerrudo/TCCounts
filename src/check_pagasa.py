"""
Checks every value reported from src/pagasa.py, recomputing what can be
recomputed from the inputs rather than reading it back.

Usage
    python src/check_pagasa.py LIST.xlsx OUTDIR data/
Exits non-zero if any check fails.
"""
import os
import sys

import numpy as np
import openpyxl
import pandas as pd

from calibrate import arma_sim, detrended, fitted_models
from classes import fit_many
from dating import Binomial, Poisson, date_series
from multiverse import fit

LST, OUT, DATA = sys.argv[1], sys.argv[2], sys.argv[3]
rd = lambda f: pd.read_csv(os.path.join(OUT, f))
FAILS, N = [], [0]


def check(label, got, want, tol=0.0):
    N[0] += 1
    ok = (got == want) if isinstance(want, str) else abs(float(got) - float(want)) <= tol + 1e-12
    print(f"{'OK      ' if ok else 'MISMATCH'} {label}: got {got}, reported {want}" + (f" (tol {tol})" if tol else ""))
    if not ok:
        FAILS.append(label)


# ---- inputs, read independently of pagasa.read_list
ws = openpyxl.load_workbook(LST, data_only=True)["List of TC"]
years, types = [], []
for r in range(5, ws.max_row + 1):
    y = ws.cell(r, 2).value
    if isinstance(y, (int, float)) and 1900 < y < 2100:
        years.append(int(y))
        types.append(str(ws.cell(r, 4).value).strip())
yr = pd.Series(years)
check("listed storms, 1960-2024", len(yr), 1250)
check("listed storms, 1960-2023", int(yr.between(1960, 2023).sum()), 1232)

S = pd.read_csv(os.path.join(OUT, "pagasa_storms.csv"), index_col=0)
A = pd.read_csv(os.path.join(OUT, "pagasa_annual.csv"), index_col=0)
U = pd.read_csv(os.path.join(OUT, "pagasa_ibtracs_par.csv"), index_col=0)
T, C, D = rd("pagasa_trends.csv"), rd("pagasa_calibrated_p.csv"), rd("pagasa_dating.csv")
log = pd.read_csv(os.path.join(OUT, "pagasa_tracks_log.csv"), index_col=0).value

cnt = yr.value_counts().reindex(range(1960, 2024), fill_value=0)
check("annual list counts equal the spreadsheet, years differing", int((cnt.values != A.pagasa_all.values).sum()), 0)
check("TS or stronger plus depression equals the list, years differing",
      int(((A.pagasa_ts + A.pagasa_td) != A.pagasa_all).sum()), 0)

# ---- linking and matching
check("listed storms linked to a track", int(S.tracks.notna().sum()), 1250)
check("linked with a fix inside the PAR", int(log["list_linked_inside_par"]), 1247)
check("tracks linked to two listed storms", int(log["tracks_linked_twice"]), 0)
SL = S[S.YEAR.between(1960, 2023)]
m = SL[SL.SID.notna()]
check("listed 1960-2023 matched to IBTrACS within 200 km", len(m), 1103)
check("matched to an IBTrACS PAR main-track storm", int(SL.ib_in_par_main.sum()), 1093)
check("SIDs used twice", int(S.SID.dropna().duplicated().sum()), 0)
check("largest accepted mean separation <= 200 km", float(m.mean_km.max() <= 200), 1)
check("median separation of matches, km", round(float(m.mean_km.median())), 39, 0)
check("unmatched listed storms, 1960-2023", int(SL.SID.isna().sum()), 129)
check("  of which listed as TD", int((SL.SID.isna() & (SL.TYPE == "TD")).sum()), 106)
check("  of which share an IBTrACS storm with another listed storm",
      int((SL.unmatched_reason == "its IBTrACS storm is matched to another listed storm").sum()), 9)

# ---- IBTrACS side against the version 4 class counts
cc = pd.read_csv(os.path.join(DATA, "class_counts.csv"))
cc = cc[cc.domain == "PAR"].set_index("SEASON").loc[1960:2023]
for k in ("total", "ts_or_stronger", "depression_only", "below_depression", "no_intensity"):
    check(f"IBTrACS {k} equals class_counts.csv, years differing", int((cc[k].values != A[f"ib_{k}"].values).sum()), 0)
n = U.groupby("ib_class").size()
lst = U.groupby("ib_class").in_list.sum()
for k, nn, ll, pct in (("ts_or_stronger", 1027, 968, 94.3), ("depression_only", 340, 122, 35.9),
                       ("below_depression", 24, 2, 8.3), ("no_intensity", 26, 1, 3.8)):
    check(f"IBTrACS PAR {k}, storms", int(n[k]), nn)
    check("  listed by PAGASA", int(lst[k]), ll)
    check("  listed share %", round(100 * lst[k] / n[k], 1), pct, 0.05)
check("IBTrACS PAR storms 1960-2023", len(U), 1417)
check("  listed by PAGASA", int(U.in_list.sum()), 1093)
dep400 = U[U.ib_class == "depression_only"].in_list_400km.mean() * 100
check("depression-only listed share at 400 km %", round(dep400, 1), 41.2, 0.05)
era = pd.cut(U.SEASON, [1959, 1979, 2009, 2023], labels=["a", "b", "c"])
for k, vals in (("ts_or_stronger", (90.8, 95.2, 97.6)), ("depression_only", (29.0, 33.8, 62.0))):
    got = U[U.ib_class == k].groupby(era[U.ib_class == k], observed=True).in_list.mean() * 100
    for e, v in zip("abc", vals):
        check(f"{k} listed share, era {e} %", round(float(got[e]), 1), v, 0.05)

# ---- crosstab of listed storms
x = SL[SL.ib_in_par_main]
check("listed, IBTrACS TS or stronger", int((x.ib_class == "ts_or_stronger").sum()), 968)
check("  PAGASA TD inside the PAR among them", int(((x.ib_class == "ts_or_stronger") & ~x.ts_6h).sum()), 47)
check("listed, IBTrACS depression only, PAGASA TS or stronger", int(((x.ib_class == "depression_only") & x.ts_6h).sum()), 12)

# ---- means
for col, vals in (("pagasa_all", (19.25, 20.25, 19.23, 17.86)), ("pagasa_ts", (14.92, 15.45, 15.0, 14.0)),
                  ("pagasa_td", (4.33, 4.8, 4.23, 3.86)), ("ib_total", (22.14, 25.65, 21.07, 19.43)),
                  ("ib_ts_or_stronger", (16.05, 16.9, 15.93, 15.07))):
    for (a, b), v in zip(((1960, 2023), (1960, 1979), (1980, 2009), (2010, 2023)), vals):
        check(f"mean {col} {a}-{b}", round(float(A[col].loc[a:b].mean()), 2), v, 0.005)

# ---- trend tests, recomputed
for name in T.series.unique():
    g = T[T.series == name].set_index("start")
    bad = 0
    for y in g.index:
        sl, lo, hi, vd = fit(A[name].loc[y:2023])
        bad += int(abs(sl - g.loc[y, "slope"]) > 1e-12 or vd != g.loc[y, "verdict"])
    check(f"trend recomputed, {name}, starts differing", bad, 0)
rep = {"pagasa_all": (-0.0535, 0.0), "pagasa_ts": (-0.0401, 0.0), "pagasa_td": (-0.0134, 0.0),
       "pagasa_ts_listed": (-0.0500, 14.6), "ib_total": (-0.1455, 51.2), "ib_ts_or_stronger": (-0.0478, 2.4)}
for name, (s0, pct) in rep.items():
    g = T[T.series == name]
    check(f"{name} slope from 1960", round(float(g[g.start == 1960].slope.iloc[0]), 4), s0, 0.00005)
    check(f"{name} starts with a significant decline %", round(100 * float((g.verdict == "-").mean()), 1), pct, 0.05)
check("pagasa_ts negative slopes %", round(100 * float((T[T.series == "pagasa_ts"].slope < 0).mean()), 1), 90.2, 0.05)
check("pagasa_all starts with a significant increase", int((T[T.series == "pagasa_all"].verdict == "+").sum()), 0)

# ---- calibrated p-values: AIC model reproduced with a fresh seed, within Monte Carlo error
for name, y0, rep_p in (("pagasa_all", 1960, 0.064), ("pagasa_ts", 1960, 0.067), ("pagasa_td", 1960, 0.39),
                        ("ib_ts_or_stronger", 1960, 0.029), ("ib_total", 1977, 0.011)):
    models = fitted_models(detrended(A[name].to_numpy(float)))
    lab = [k for k in models if "AIC" in k][0]
    ar, ma, sig = models[lab]
    vv = A[name].loc[y0:2023].to_numpy(float)
    xx = np.arange(len(vv), dtype=float)
    sl = fit_many(vv[None, :], xx)[0][0]
    R = 20000
    Z = arma_sim(ar, ma, sig, len(vv), R, np.random.default_rng(77 + y0))
    p = float(np.mean(np.abs(fit_many(Z, xx)[0]) >= abs(sl)))
    tol = 4 * np.sqrt(max(p, 1e-3) * (1 - p) / R) + 0.0015
    check(f"calibrated p, {name} from {y0}, {lab}, fresh seed", round(p, 3), rep_p, tol)
check("ARMA(1,1) boundary fits left blank", int(C[C.boundary_fit].p_two_sided.notna().sum()), 0)
check("ib_total from 1960, p below 0.001 (AIC)",
      float(C[(C.series == "ib_total") & (C.start == 1960) & (C.model_role == "AIC")].p_two_sided.iloc[0] < 0.001), 1)

# ---- step models, recomputed
t = np.arange(1960, 2024)
for name, sel in (("pagasa_all", "constant"), ("pagasa_ts", "constant"), ("pagasa_td", "constant"),
                  ("ib_ts_or_stronger", "constant"), ("ib_total", "smooth")):
    r = date_series(Poisson(A[name].to_numpy(float)), t)
    check(f"dating {name}", r["selected"], sel)
r = date_series(Binomial(A.ib_listed.to_numpy(float), A.ib_total.to_numpy(float)), t)
check("dating listed share of IBTrACS total, breaks", r["breaks"], "1974 2002")
check("  segments", r["segments"], "1960-1973: 64.1%; 1974-2001: 77.7%; 2002-2023: 87.9%")
check("  support", r["support"], "1974-1977 1999-2003")

print(f"\n{N[0]} checks, {len(FAILS)} failed")
sys.exit(1 if FAILS else 0)
